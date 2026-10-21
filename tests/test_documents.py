import unittest
import tempfile
from pathlib import Path
from unittest.mock import Mock, patch
from core.documents import parse_request, create_from_request, render_document

class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)

    def test_formats_and_unrelated_commands(self):
        for text, kind in [('V, crie um PDF sobre física', 'pdf'), ('gere uma planilha de gastos', 'xlsx'), ('faça um documento sobre história','docx')]:
            self.assertEqual(parse_request(text)[0], kind)
        self.assertIsNone(parse_request('o que é um PDF?'))
        self.assertIsNone(parse_request('não crie um PDF'))

    def test_pdf_and_word_round_trip_and_unique_files(self):
        from docx import Document
        from pypdf import PdfReader
        data = {'title':'Revisão & planejamento', 'paragraphs':['Ação, ciência e educação. <teste>', 'Segundo parágrafo.']}
        for kind in ('pdf','docx'):
            with self.subTest(kind=kind):
                first = render_document(kind,data,self.folder)
                second = render_document(kind,data,self.folder)
                self.assertNotEqual(first,second)
                if kind == 'pdf':
                    text = ''.join(page.extract_text() for page in PdfReader(first).pages)
                else:
                    text = '\n'.join(p.text for p in Document(first).paragraphs)
                self.assertIn('educação',text)
                self.assertIn('<teste>',text)
        self.assertFalse(list(self.folder.glob('.partial-*')))

    def test_sheet_numbers_literal_formulas_and_layout(self):
        from openpyxl import load_workbook
        p = render_document('xlsx', {'title':'Orçamento', 'columns':['Item','Valor'], 'rows':[['Aluguel',1500],['=HYPERLINK("evil")',20.5]]}, self.folder)
        book=load_workbook(p); sheet=book.active
        self.assertEqual(sheet['B2'].value,1500)
        self.assertEqual(sheet['A3'].data_type,'s')
        self.assertEqual(sheet.freeze_panes,'A2')
        self.assertEqual(sheet.auto_filter.ref,'A1:B3')
        book.close()

    def test_failed_content_does_not_leave_file(self):
        brain=Mock(); brain.chat.return_value='não é JSON'
        with self.assertRaises(ValueError):
            create_from_request('pdf','física',brain,directory=self.folder)
        with self.assertRaises(ValueError):
            render_document('xlsx', {'title':'Teste','columns':['A'], 'rows':[[1,2]]}, self.folder)
        self.assertFalse(list(self.folder.iterdir()))

    def test_offline_literal_and_ai_content(self):
        create_from_request('docx',': Texto fornecido pelo usuário.',directory=self.folder)
        brain=Mock(); brain.chat.return_value='```json\n{"title":"Estudos","paragraphs":["Revisar notas."]}\n```'
        create_from_request('pdf','plano de estudos',brain,directory=self.folder)
        self.assertEqual(len(list(self.folder.iterdir())),2)

    def test_engine_routes_documents_before_other_actions(self):
        from core.functional_engine import FunctionalEngine
        engine=FunctionalEngine.__new__(FunctionalEngine)
        engine.brain=None; engine.history=[]; engine.projects=Mock()
        engine._append_history=Mock()
        with patch('core.documents.create_from_request',return_value=self.folder/'arquivo.pdf') as create:
            result=engine.handle('crie um PDF: exemplo')[0]
        self.assertTrue(result.ok)
        self.assertEqual(result.data['kind'],'document')
        create.assert_called_once()

if __name__=='__main__': unittest.main()

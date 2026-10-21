"""Local document creation. Content is data, never executable code."""
from pathlib import Path
from datetime import datetime
import json
import math
import re
import uuid
import unicodedata

OUTPUT_DIR = Path(__file__).resolve().parents[1] / 'exports'

def parse_request(text):
    raw = re.sub(r'^\s*(?:v|vi|jarvis)\b[ ,:]*', '', text, flags=re.I)
    match = re.match(r'^(?:crie|cria|criar|gere|gerar|faca|faça|monte|exporte|salve)\s+(?:(?:um|uma|o|a)\s+)?(pdf|planilha|excel|xlsx|documento|word|docx)\b(.*)$', raw, re.I | re.S)
    if not match:
        return None
    kind, detail = match.groups()
    kind = {'planilha':'xlsx', 'excel':'xlsx', 'documento':'docx', 'word':'docx'}.get(kind.lower(), kind.lower())
    detail = re.sub(r'^\s*(?:sobre|de|com|em)\s+', '', detail.strip(), flags=re.I)
    return kind, detail

def validate_content(value, kind):
    if not isinstance(value, dict):
        raise ValueError('A resposta precisa ser um objeto JSON.')
    title = value.get('title', '')
    if not isinstance(title, str) or not title.strip() or len(title) > 160:
        raise ValueError('Título ausente ou muito longo.')
    if kind == 'xlsx':
        columns, rows = value.get('columns'), value.get('rows')
        if not isinstance(columns, list) or not 1 <= len(columns) <= 30 or not all(isinstance(c, str) and 0 < len(c) <= 150 for c in columns):
            raise ValueError('Informe de 1 a 30 colunas com nomes válidos.')
        if not isinstance(rows, list) or len(rows) > 1000:
            raise ValueError('A planilha aceita até 1.000 linhas por pedido.')
        for row in rows:
            if not isinstance(row, list) or len(row) != len(columns):
                raise ValueError('Cada linha deve ter o mesmo número de células que o cabeçalho.')
            for cell in row:
                if cell is not None and type(cell) not in (str, int, float, bool):
                    raise ValueError('Valor de célula inválido.')
                if isinstance(cell, str) and len(cell) > 4000:
                    raise ValueError('Texto da célula muito longo.')
                if isinstance(cell, float) and not math.isfinite(cell):
                    raise ValueError('Número inválido.')
    else:
        paragraphs = value.get('paragraphs')
        if not isinstance(paragraphs, list) or not paragraphs or len(paragraphs) > 300 or not all(isinstance(p, str) and len(p) <= 12000 for p in paragraphs):
            raise ValueError('O documento precisa conter parágrafos de texto válidos.')
        if not any(p.strip() for p in paragraphs):
            raise ValueError('O documento está vazio.')
    if len(json.dumps(value, ensure_ascii=False)) > 150000:
        raise ValueError('Conteúdo muito grande; divida o pedido.')
    return value

def render_document(kind, content, directory=None):
    validate_content(content, kind)
    directory = Path(directory or OUTPUT_DIR)
    directory.mkdir(parents=True, exist_ok=True)
    stem = unicodedata.normalize('NFKD', content['title']).encode('ascii', 'ignore').decode()
    stem = re.sub(r'[^a-zA-Z0-9_-]+', '-', stem).strip('-')[:60] or 'documento'
    destination = directory / f"{datetime.now():%Y%m%d-%H%M%S}-{stem}-{uuid.uuid4().hex[:8]}.{kind}"
    temporary = destination.with_name('.partial-' + destination.name)
    try:
        if kind == 'xlsx':
            from openpyxl import Workbook, load_workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
            book = Workbook(); sheet = book.active; sheet.title = 'Dados'
            for row in [content['columns']] + content['rows']:
                sheet.append(row)
                # Spreadsheet content must never become an executable formula.
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='8B294B')
            sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
            for i, name in enumerate(content['columns'], 1):
                sheet.column_dimensions[get_column_letter(i)].width = min(45, max(18, len(name) + 4))
            book.save(temporary); book.close()
            check = load_workbook(temporary); check.close()
        elif kind == 'docx':
            from docx import Document
            from docx.shared import Pt, Inches, RGBColor
            doc = Document()
            section = doc.sections[0]
            section.top_margin = section.bottom_margin = Inches(.8)
            style = doc.styles['Normal']; style.font.name = 'Calibri'; style.font.size = Pt(11)
            style.paragraph_format.space_after = Pt(8)
            doc.add_heading(content['title'], 0)
            for text in content['paragraphs']:
                doc.add_paragraph(text)
            doc.core_properties.title = content['title']
            doc.save(temporary); Document(temporary)
        elif kind == 'pdf':
            from html import escape
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from reportlab.lib import colors
            font_path = Path('C:/Windows/Fonts/arial.ttf')
            font_name = 'Helvetica'
            if font_path.is_file():
                if 'VArial' not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont('VArial', str(font_path)))
                font_name = 'VArial'
            styles = getSampleStyleSheet()
            for style in (styles['Title'], styles['BodyText']):
                style.fontName = font_name
            styles['Title'].textColor = colors.HexColor('#8b294b')
            styles['BodyText'].leading = 16
            story = [Paragraph(escape(content['title']), styles['Title']), Spacer(1, 16)]
            for text in content['paragraphs']:
                story.extend([Paragraph(escape(text).replace('\n', '<br/>'), styles['BodyText']), Spacer(1, 9)])
            SimpleDocTemplate(str(temporary), title=content['title']).build(story)
            from pypdf import PdfReader
            if not PdfReader(temporary).pages:
                raise ValueError('PDF gerado sem páginas.')
        else:
            raise ValueError('Formato não suportado.')
        temporary.replace(destination)
        return destination.resolve()
    finally:
        temporary.unlink(missing_ok=True)

def create_from_request(kind, detail, brain=None, history=None, directory=None):
    if not detail:
        raise ValueError('Diga o assunto ou forneça o conteúdo. Exemplo: crie um PDF sobre meu plano de estudos.')
    # Explicit literal content and structured spreadsheets work without an AI provider.
    if detail.startswith(':') and kind != 'xlsx':
        text = detail[1:].strip()
        content = {'title': 'Documento do V', 'paragraphs': text.split('\n\n')}
    elif detail.startswith('{'):
        content = json.loads(detail)
    else:
        if brain is None:
            raise ValueError('Conecte uma IA para elaborar o conteúdo, ou envie texto após dois-pontos.')
        schema = '{"title":"Título", "columns":["Coluna"], "rows":[["Valor"]]}' if kind == 'xlsx' else '{"title":"Título", "paragraphs":["Parágrafo"]}'
        prompt = ('Prepare o conteúdo de um arquivo. Responda SOMENTE JSON válido no formato ' + schema +
                  '. Use português do Brasil. Não execute ações. Não invente dados pessoais, valores reais, fontes ou fatos atuais; '
                  'quando faltarem dados, crie um modelo com campos vazios ou exemplos claramente identificados. '
                  'Não escreva fórmulas de planilha. Pedido do usuário: ' + detail)
        answer = brain.chat(prompt, history or [], {})
        if not isinstance(answer, str) or not answer.strip():
            raise ValueError('A IA não retornou conteúdo. Verifique a conexão e tente novamente.')
        answer = re.sub(r'^```(?:json)?\s*|\s*```$', '', answer.strip())
        content = json.loads(answer)
    return render_document(kind, content, directory)
